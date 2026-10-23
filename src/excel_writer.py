"""Escritura del Excel con formato listo para copiar y pegar."""
from __future__ import annotations

from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .buckets import (ADEL, CIERRE, TOTAL, Columna, Params, descripcion_ciclo,
                      etiqueta_periodo)

F = "Arial"
AZUL_OSC, AZUL_CLARO = "1F3864", "DDEBF7"
borde_c = Side(style="thin", color="8EA9DB")
BORDE = Border(left=borde_c, right=borde_c, top=borde_c, bottom=borde_c)
ROJO = "FF0000"


def _fill(hexa):  # noqa: D401
    return PatternFill("solid", start_color=hexa, end_color=hexa)


def _num(v):
    """Ceros/NaN -> celda vacía."""
    if v is None or pd.isna(v) or v == 0:
        return None
    return float(v)


# ---------------------------------------------------------------- RES MES
def escribir_res_mes(wb: Workbook, tabla: pd.DataFrame) -> None:
    ws = wb.active
    ws.title = "RES MES"
    ws["A1"], ws["B1"] = "Suma de Real Pago Dolar", "Etiquetas de columna"
    for c in ("A1", "B1"):
        ws[c].font = Font(name=F, bold=True, size=10)
        ws[c].fill = _fill(AZUL_CLARO)

    ws["A2"] = "Etiquetas de fila"
    for j, nombre in enumerate(tabla.columns, start=2):
        ws.cell(2, j, nombre)
    for j in range(1, len(tabla.columns) + 2):
        c = ws.cell(2, j)
        c.font = Font(name=F, bold=True, size=10)
        c.fill = _fill(AZUL_CLARO)
        c.alignment = Alignment(wrap_text=True, vertical="bottom",
                                horizontal="left" if j == 1 else "right")
        c.border = BORDE
    ws.row_dimensions[2].height = 42

    for i, (idx, fila) in enumerate(tabla.iterrows(), start=3):
        es_total = isinstance(idx, str)
        c0 = ws.cell(i, 1, idx.to_pydatetime() if not es_total else idx)
        if not es_total:
            c0.number_format = "d/m/yyyy"
        c0.alignment = Alignment(horizontal="left")
        for j, v in enumerate(fila, start=2):
            c = ws.cell(i, j, _num(v))
            c.number_format = "0"
        for j in range(1, len(tabla.columns) + 2):
            c = ws.cell(i, j)
            c.font = Font(name=F, size=10, bold=es_total or j == len(tabla.columns) + 1)
            c.border = BORDE
            if es_total:
                c.fill = _fill(AZUL_CLARO)

    ws.column_dimensions["A"].width = 18
    for j in range(2, len(tabla.columns) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 15
    ws.freeze_panes = "B3"


# ---------------------------------------------------------------- hoja por asesora
def _encabezado(ws, fila: int, cols: list[Columna], etiqueta: str) -> None:
    ws.row_dimensions[fila].height = 42
    for j in (1, 2):
        ws.cell(fila, j).fill = _fill(AZUL_CLARO)
        ws.cell(fila, j).border = BORDE
    c = ws.cell(fila, 2, etiqueta)
    c.font = Font(name=F, bold=True, size=10, color="1F3864")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for k, col in enumerate(cols, start=3):
        c = ws.cell(fila, k, col.label)
        oscuro = col.estilo == "anio"
        c.fill = _fill(AZUL_OSC if oscuro else AZUL_CLARO)
        c.font = Font(name=F, bold=True, size=10, color="FFFFFF" if oscuro else "1F3864")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE


def _fila_valores(ws, fila: int, serie: pd.Series, cols: list[Columna], negrita=False, relleno=None) -> None:
    for k, col in enumerate(cols, start=3):
        c = ws.cell(fila, k, _num(serie[col.key]))
        c.number_format = "0"
        c.font = Font(name=F, size=10, bold=negrita or col.key == TOTAL)
        c.border = BORDE
        c.alignment = Alignment(horizontal="right")
        if relleno:
            c.fill = _fill(relleno)


def _titulo(ws, fila: int, texto: str) -> None:
    c = ws.cell(fila, 2, texto)
    c.font = Font(name=F, italic=True, size=9, color="7F7F7F")


def escribir_hoja_asesora(wb: Workbook, nombre_hoja: str, titulo: str, p: Params,
                          cols: list[Columna], dias: pd.DatetimeIndex,
                          t_monto: pd.DataFrame, t_cuotas: pd.DataFrame,
                          per_monto: pd.DataFrame, per_cuotas: pd.DataFrame,
                          periodos: list, etiqueta: str) -> None:
    ws = wb.create_sheet(nombre_hoja)
    ws["B1"] = titulo
    ws["B1"].font = Font(name=F, bold=True, size=11, color="0000FF")

    # ---- Tabla 1: por día (Monto / Numero de cuotas)
    fila = 3
    _titulo(ws, fila, "Tabla 1 · por Fecha Pago (Monto = suma, Nº de cuotas = conteo de Real Pago Dolar)")
    fila += 1
    _encabezado(ws, fila, cols, etiqueta)
    fila += 1
    for d in dias:
        rojo = d.weekday() >= 5  # sábado/domingo
        for etiq, tabla in (("Monto", t_monto), ("Numero de cuotas", t_cuotas)):
            c = ws.cell(fila, 1, d.to_pydatetime())
            c.number_format = "dd/mm"
            c.alignment = Alignment(horizontal="center")
            c.font = Font(name=F, size=10, color=ROJO if rojo else "000000")
            c.border = BORDE
            b = ws.cell(fila, 2, etiq)
            b.font = Font(name=F, size=10)
            b.border = BORDE
            _fila_valores(ws, fila, tabla.loc[d], cols)
            fila += 1

    # ---- Tablas 2 y 3: por PERIODO
    for titulo_t, tm in (("Tabla 2 · por PERIODO · Monto (suma de Real Pago Dolar)", per_monto),
                         ("Tabla 3 · por PERIODO · Numero de cuotas (conteo de Real Pago Dolar)", per_cuotas)):
        fila += 1
        _titulo(ws, fila, titulo_t)
        fila += 1
        _encabezado(ws, fila, cols, etiqueta)
        fila += 1
        ws.cell(fila, 1).border = BORDE
        b = ws.cell(fila, 2, "TOTAL")
        b.font = Font(name=F, bold=True, size=10)
        b.border = BORDE
        _fila_valores(ws, fila, tm.sum(axis=0), cols, negrita=True, relleno=AZUL_CLARO)
        fila += 1
        for per in periodos:
            ws.cell(fila, 1).border = BORDE
            b = ws.cell(fila, 2, etiqueta_periodo(per))
            b.font = Font(name=F, size=10, color="7F7F7F")
            b.alignment = Alignment(horizontal="left", indent=1)
            b.border = BORDE
            _fila_valores(ws, fila, tm.loc[per], cols)
            fila += 1

    # ---- Nota (fuera de las tablas)
    fila += 1
    ws.cell(fila, 2, descripcion_ciclo(p)).font = Font(name=F, italic=True, size=9, color="7F7F7F")
    ws.cell(fila + 1, 2, f"Generado: {datetime.now():%d/%m/%Y %H:%M}").font = Font(
        name=F, italic=True, size=9, color="7F7F7F")

    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 26
    for k, col in enumerate(cols, start=3):
        ws.column_dimensions[get_column_letter(k)].width = 12 if col.key in (ADEL, CIERRE, TOTAL) else 9.5
